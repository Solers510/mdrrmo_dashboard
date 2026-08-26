import unittest
from datetime import datetime, timezone
from io import BytesIO

from openpyxl import load_workbook

from services.export_service import (
    build_excel_report,
    build_pdf_report,
    report_filename,
)
from services.report_service import (
    REPORT_MODE_OFFICIAL,
    REPORT_MODE_PROVISIONAL,
    _select_reconciliation_payload,
    canonical_snapshot_json,
    snapshot_sha256,
)


SAMPLE_RECORD = {
    "id": 7,
    "event_id": 1,
    "generation_key": "11111111-1111-1111-1111-111111111111",
    "report_type": "Situation Report",
    "report_mode": "Official Validated",
    "sitrep_number": "3",
    "generated_by": "Test User - Administrator",
    "generated_at": datetime(
        2026,
        8,
        9,
        12,
        0,
        tzinfo=timezone.utc,
    ),
    "snapshot_sha256": "a" * 64,
    "snapshot_json": {
        "schema_version": 1,
        "event": {
            "event_name": "Luis",
            "classification": "Tropical Depression",
            "hazard_type": "Tropical Cyclone",
            "alert_code": "BLUE",
            "eoc_status": "Monitoring",
            "official_reference": "TEST-REF",
            "situation_overview": "Test overview.",
        },
        "population_summary": {
            "affected_barangays": 1,
            "affected_families": 10,
            "affected_individuals": 40,
            "inside_ec_families": 5,
            "inside_ec_individuals": 20,
            "outside_ec_families": 2,
            "outside_ec_individuals": 8,
            "displaced_families": 7,
            "displaced_individuals": 28,
            "affected_not_displaced_families": 3,
            "affected_not_displaced_individuals": 12,
            "pending_rescue_requests": 1,
            "impassable_roads": 0,
            "interrupted_power": 1,
            "interrupted_water": 0,
            "reports_received": 1,
            "missing_reports": 29,
            "pending_validation": 0,
            "needs_correction": 0,
        },
        "barangays": [
            {
                "barangay_name": "Test Barangay",
                "situation_status": "Affected",
                "affected_families": 10,
                "affected_individuals": 40,
                "inside_ec_families": 5,
                "inside_ec_individuals": 20,
                "outside_ec_families": 2,
                "outside_ec_individuals": 8,
                "flood_status": "No Flooding",
                "flood_depth_cm": 0,
                "road_status": "Passable",
                "power_status": "Interrupted",
                "water_status": "Normal",
                "rescue_requests": 1,
                "validation_status": "Validated",
                "source": "Test source",
                "recorded_at": "2026-08-09T12:00:00+08:00",
                "remarks": "",
            }
        ],
        "evacuation_summary": {
            "open_centers": 1,
            "families": 5,
            "individuals": 20,
            "safe_capacity": 50,
            "over_capacity_centers": 0,
            "critical_food": 0,
            "critical_water": 0,
            "medical_cases": 1,
        },
        "evacuation_centers": [
            {
                "center_name": "Test EC",
                "barangay_name": "Test Barangay",
                "status": "Open",
                "families": 5,
                "individuals": 20,
                "safe_capacity": 50,
                "children": 4,
                "senior_citizens": 2,
                "pwd": 1,
                "pregnant_women": 1,
                "medical_cases": 1,
                "food_status": "Sufficient",
                "water_status": "Sufficient",
                "electricity_status": "Available",
                "sanitation_status": "Functional",
                "validation_status": "Validated",
                "source": "Camp manager",
                "recorded_at": "2026-08-09T12:00:00+08:00",
                "remarks": "",
            }
        ],
        "reconciliation_summary": {
            "match": 1,
            "mismatch": 0,
            "missing_source": 0,
            "allocation_conflict": 0,
        },
        "reconciliation": [],
        "incidents": [],
    },
}


class SnapshotIntegrityTests(unittest.TestCase):
    def test_canonical_json_is_order_independent(self):
        first = {
            "b": 2,
            "a": 1,
        }
        second = {
            "a": 1,
            "b": 2,
        }

        self.assertEqual(
            canonical_snapshot_json(first),
            canonical_snapshot_json(second),
        )
        self.assertEqual(
            snapshot_sha256(first),
            snapshot_sha256(second),
        )

    def test_hash_changes_when_snapshot_changes(self):
        first = {"value": 1}
        second = {"value": 2}

        self.assertNotEqual(
            snapshot_sha256(first),
            snapshot_sha256(second),
        )

    def test_reconciliation_payload_follows_report_mode(self):
        dashboard = {
            "provisional_reconciliation_summary": {"match": 2},
            "provisional_reconciliation_rows": [{"mode": "provisional"}],
            "provisional_reconciliation_available": True,
            "official_reconciliation_summary": {"match": 1},
            "official_reconciliation_rows": [{"mode": "official"}],
            "official_reconciliation_available": False,
        }

        provisional = _select_reconciliation_payload(
            dashboard,
            REPORT_MODE_PROVISIONAL,
        )
        official = _select_reconciliation_payload(
            dashboard,
            REPORT_MODE_OFFICIAL,
        )

        self.assertEqual(provisional[0], {"match": 2})
        self.assertEqual(provisional[1], [{"mode": "provisional"}])
        self.assertTrue(provisional[2])
        self.assertEqual(official[0], {"match": 1})
        self.assertEqual(official[1], [{"mode": "official"}])
        self.assertFalse(official[2])


class ExportTests(unittest.TestCase):
    def test_excel_export_contains_expected_sheets(self):
        data = build_excel_report(
            SAMPLE_RECORD
        )

        self.assertTrue(
            data.startswith(b"PK")
        )

        workbook = load_workbook(
            BytesIO(data),
            read_only=True,
        )

        self.assertEqual(
            workbook.sheetnames,
            [
                "Executive Summary",
                "Barangay Situation",
                "Evacuation Centers",
                "Reconciliation",
                "Incident Log",
                "Snapshot Metadata",
            ],
        )

        summary = workbook[
            "Executive Summary"
        ]
        self.assertEqual(
            summary["A1"].value,
            "MDRRMO Naic Situation Report",
        )
        labels = {
            summary.cell(row=row, column=1).value
            for row in range(1, summary.max_row + 1)
        }
        self.assertIn("Barangays with Affected People", labels)
        self.assertNotIn("Affected, Not Displaced - Families", labels)

    def test_pdf_export_is_pdf(self):
        data = build_pdf_report(
            SAMPLE_RECORD
        )

        self.assertTrue(
            data.startswith(b"%PDF-")
        )
        self.assertGreater(
            len(data),
            1000,
        )

    def test_filename_is_filesystem_safe(self):
        name = report_filename(
            SAMPLE_RECORD,
            extension="pdf",
        )

        self.assertTrue(
            name.endswith(".pdf")
        )
        self.assertNotIn(" ", name)


if __name__ == "__main__":
    unittest.main()
